"""
Editable Word (.docx) Report Generator for OIML R 76 NAWI Test Reports
Generates standardized DOCX documents for official government records and editing.
"""

import io
from datetime import datetime
from typing import Dict, Any
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


def generate_nawi_docx_report(session_data: Dict[str, Any], instrument_data: Dict[str, Any], manufacturer_data: Dict[str, Any], evaluation_result: Dict[str, Any]) -> bytes:
    """
    Generates an official Word (.docx) report for NAWI evaluation.
    """
    doc = Document()

    # Page setup - 0.75 inch margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Header Paragraphs
    p_gov = doc.add_paragraph()
    p_gov.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_gov = p_gov.add_run("GOVERNMENT OF INDIA\nMINISTRY OF CONSUMER AFFAIRS, FOOD & PUBLIC DISTRIBUTION\nDEPARTMENT OF CONSUMER AFFAIRS — LEGAL METROLOGY DIVISION")
    run_gov.font.name = "Arial"
    run_gov.font.size = Pt(11)
    run_gov.font.bold = True
    run_gov.font.color.rgb = RGBColor(30, 58, 138)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("OIML R 76-1 TYPE EVALUATION TEST REPORT")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(14)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(6, 95, 70)

    # Info Table
    info_table = doc.add_table(rows=6, cols=4)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.style = 'Table Grid'

    reviewer_status = session_data.get("status") or evaluation_result.get("overall_status", "PENDING")
    reviewer_name = session_data.get("reviewer_name", "Official Reviewer")
    testing_officer = session_data.get("testing_officer", "Official Technician")

    fields = [
        ("Report Code:", session_data.get("session_code", "N/A"), "Test Date:", str(session_data.get("test_date", datetime.now().strftime("%Y-%m-%d")))),
        ("Manufacturer:", manufacturer_data.get("name", "N/A"), "Country:", manufacturer_data.get("country", "India")),
        ("Model Name:", instrument_data.get("model_name", "N/A"), "Serial No:", instrument_data.get("serial_number", "N/A")),
        ("Accuracy Class:", instrument_data.get("accuracy_class", "CLASS_III"), "Max Cap (Max):", f"{instrument_data.get('max_capacity', 0)} {instrument_data.get('unit', 'kg')}"),
        ("Verification Interval (e):", f"{instrument_data.get('verification_scale_interval', 0)} {instrument_data.get('unit', 'kg')}", "Testing Officer:", testing_officer),
        ("Reviewer Name:", reviewer_name, "Reviewer Action:", reviewer_status)
    ]

    for r_idx, row_vals in enumerate(fields):
        row = info_table.rows[r_idx]
        for c_idx in range(4):
            cell = row.cells[c_idx]
            cell.text = row_vals[c_idx]
            if c_idx % 2 == 0:
                shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F3F4F6"/>')
                cell._tc.get_or_add_tcPr().append(shading_elm)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 1. Environmental Conditions
    h1 = doc.add_heading("1. Environmental Conditions", level=2)
    h1.runs[0].font.color.rgb = RGBColor(30, 58, 138)

    p_env = doc.add_paragraph(
        f"Temperature: {session_data.get('temperature_c', 20.0)} °C  |  "
        f"Relative Humidity: {session_data.get('humidity_percent', 50.0)} %  |  "
        f"Pressure: {session_data.get('pressure_hpa', 1013.25)} hPa"
    )

    # 2. Weighing Test Table
    h2 = doc.add_heading("2. Weighing Span & Linearity Results", level=2)
    h2.runs[0].font.color.rgb = RGBColor(30, 58, 138)

    test_results = evaluation_result.get("test_results", {})
    span_points = test_results.get("span_test", {}).get("points", [])

    if span_points:
        span_table = doc.add_table(rows=len(span_points) + 1, cols=5)
        span_table.style = 'Table Grid'
        headers = ["Test Load", "Indication", "Error E", "MPE", "Status"]
        hdr_row = span_table.rows[0]
        for idx, h_text in enumerate(headers):
            cell = hdr_row.cells[idx]
            cell.text = h_text
            shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1E3A8A"/>')
            cell._tc.get_or_add_tcPr().append(shading_elm)
            p = cell.paragraphs[0]
            p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
            p.runs[0].font.bold = True

        for row_idx, pt in enumerate(span_points, start=1):
            row = span_table.rows[row_idx]
            row.cells[0].text = str(pt["load"])
            row.cells[1].text = str(pt["indication"])
            row.cells[2].text = f"{pt['error']:+.4f}"
            row.cells[3].text = f"±{pt['mpe']:.4f}"
            row.cells[4].text = pt["status"]

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 3. Signatures & Endorsement
    h3 = doc.add_heading("3. Endorsement & Signatures", level=2)
    h3.runs[0].font.color.rgb = RGBColor(30, 58, 138)

    p_endorse = doc.add_paragraph()
    p_endorse.add_run("REVIEWER ACTION & ENDORSEMENT:\n").bold = True
    run_action = p_endorse.add_run(f"This report has been officially {reviewer_status} by Reviewing Officer {reviewer_name} in accordance with OIML R 76-1:2006 guidelines.\n")
    if reviewer_status == "APPROVED":
        run_action.font.color.rgb = RGBColor(4, 120, 87)
    elif reviewer_status == "REJECTED":
        run_action.font.color.rgb = RGBColor(190, 18, 60)
    run_action.font.bold = True

    sig_p = doc.add_paragraph()
    sig_p.add_run(f"\nTested By: ___________________ ({testing_officer})\n")
    sig_p.add_run(f"Reviewed & Action Taken By: ___________________ ({reviewer_name})\n")
    run_stamp = sig_p.add_run(f"Official Registry Status: [ {reviewer_status} BY LEGAL METROLOGY DIVISION ]")
    run_stamp.font.bold = True
    if reviewer_status == "APPROVED":
        run_stamp.font.color.rgb = RGBColor(4, 120, 87)
    elif reviewer_status == "REJECTED":
        run_stamp.font.color.rgb = RGBColor(190, 18, 60)

    # Save to BytesIO
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
